"""Generate docs/Auravance-Architecture.docx from the live repos, not the resume slide."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BACKEND_DIR = Path(__file__).resolve().parents[1]
OUTPUT = BACKEND_DIR / "docs" / "Auravance-Architecture.docx"
FIGURES = BACKEND_DIR / "docs" / "architecture-diagrams"

NAVY = RGBColor(0x0B, 0x2E, 0x4A)
TEAL = RGBColor(0x0D, 0x6E, 0x7A)
MUTED = RGBColor(0x44, 0x55, 0x66)


def set_run(run, *, bold=False, size=11, color=None, italic=False):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    if color is not None:
        run.font.color.rgb = color


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = NAVY if level == 1 else TEAL
        run.font.name = "Calibri"
    return p


def add_para(doc, text, *, bold=False, italic=False, size=11, space_after=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run(text)
    set_run(run, bold=bold, italic=italic, size=size, color=MUTED if italic and not bold else None)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(item, style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        for run in p.runs:
            run.font.name = "Calibri"
            run.font.size = Pt(11)


def add_figure(doc, filename: str, caption: str) -> None:
    path = FIGURES / filename
    if not path.exists():
        raise FileNotFoundError(f"Missing diagram {path}. Run scripts/generate_architecture_diagrams.py first.")
    picture = doc.add_paragraph()
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.paragraph_format.space_before = Pt(8)
    picture.paragraph_format.space_after = Pt(2)
    run = picture.add_run()
    run.add_picture(str(path), width=Inches(6.5))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(12)
    r = cap.add_run(caption)
    set_run(r, italic=True, size=10, color=MUTED)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                set_run(run, bold=True, size=10, color=NAVY)
    for r_index, row in enumerate(rows, start=1):
        for c_index, value in enumerate(row):
            cell = table.rows[r_index].cells[c_index]
            cell.text = str(value)
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    set_run(run, size=10)
    doc.add_paragraph()
    return table


def build() -> Path:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.9)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run("Auravance")
    set_run(run, bold=True, size=28, color=NAVY)

    subtitle = doc.add_paragraph()
    run = subtitle.add_run("Industrial Predictive Intelligence Platform")
    set_run(run, size=16, color=TEAL)

    meta = doc.add_paragraph()
    run = meta.add_run(
        "Full-stack architecture: backend, frontend, ingestion, storage, "
        "models, RAG, Docker, and Terraform.\n"
        "This document describes what the repositories actually implement, "
        "and separately what remains a production target."
    )
    set_run(run, size=11, italic=True, color=MUTED)

    add_para(doc, "Repositories: auravance-backend (FastAPI) and auravance-frontend (React 18 + Vite). Date: 21 October 2026.")
    add_heading(doc, "List of figures", 2)
    add_bullets(
        doc,
        [
            "Figure 1. System context — two repositories, one HTTP contract",
            "Figure 2. Local runtime — laptop processes and data stores",
            "Figure 3. Frontend console modules",
            "Figure 4. Ingestion pipeline (REST and MQTT)",
            "Figure 5. Scoring path (Isolation Forest, XGBoost/HGB, LSTM, rules)",
            "Figure 6. RAG root-cause analysis",
            "Figure 7. SQLite demo store",
            "Figure 8. GCP demo deploy (Docker, Terraform, Firebase)",
            "Figure 9. Implemented stack versus production target",
            "Figure 10. Dashboard request path after login",
        ],
    )

    add_heading(doc, "1. Purpose of this document", 1)
    add_para(
        doc,
        "Auravance is split across two git repositories. The backend owns telemetry ingest, "
        "scoring, alerts, master data, and RAG diagnostics. The frontend is a single-page "
        "React console that calls the backend over HTTP. Local development uses SQLite and "
        "a Vite proxy. GCP demo deploy uses Cloud Run and Firebase Hosting.",
    )
    add_para(
        doc,
        "Resume-style capability statements are treated as the product north star. They are "
        "the build goal. They are not copied into this document as a description of the "
        "current tree unless matching code exists.",
    )

    add_heading(doc, "2. Product north star (target capability)", 1)
    add_para(doc, "Use this wording only as the target. Do not treat it as a description of the current tree unless the matching implementation exists.", italic=True)
    add_bullets(
        doc,
        [
            "Built a full-stack AI platform (React, FastAPI) for real-time equipment monitoring and failure prediction using XGBoost and LSTM, improving detection accuracy by 30%.",
            "Engineered scalable event-driven IoT ingestion pipelines using MQTT and Python for high-frequency sensor telemetry with low latency and high reliability.",
            "Designed scalable ETL and event-driven MQTT pipelines processing 1M+ sensor records/day, enabling predictive maintenance, anomaly detection, and RUL insights with sub-minute latency.",
            "Built a RAG-based root cause analysis system over equipment manuals to generate explainable diagnostics and maintenance recommendations.",
        ],
    )

    add_heading(doc, "3. Implemented today versus still pending", 1)
    add_table(
        doc,
        ["Capability", "In the repos today", "Still pending"],
        [
            [
                "Full-stack monitoring",
                "FastAPI /api/v1; React dashboard in auravance-frontend",
                "Production hardening (auth tenancy, HA, observability)",
            ],
            [
                "Failure / RUL models",
                "Isolation Forest anomalies; XGBoost tabular if artifacts exist else HistGradientBoosting; LSTM sequence beside tabular; rules fallback",
                "Production model serving, drift monitors, scheduled retraining",
            ],
            [
                "MQTT IoT ingest",
                "REST ingest plus a local MQTT subscriber (embedded broker by default) on plants/{plant_id}/lines/{line_id}/assets/{asset_id}/telemetry",
                "External production MQTT cluster",
            ],
            [
                "ETL / volume",
                "Offline scripts over Kaggle pump, MetroPT-3, NASA/IMS; SQLite demo DB; MQTT/REST event path into telemetry_service",
                "1M+ records/day, sub-minute latency at plant scale, TimescaleDB",
            ],
            [
                "RAG diagnostics",
                "OpenAI embeddings + TF-IDF over manuals/guides; playbook fallback",
                "pgvector index",
            ],
            [
                "Accuracy claim",
                "Leave-one-failure-out metrics in model JSON; scripts/evaluate_detection_delta.py vs rules",
                "A product-wide measured 30% detection-accuracy claim",
            ],
            [
                "Deploy",
                "Dockerfile + Terraform for Cloud Run, Artifact Registry, Secret Manager, Firebase Hosting site",
                "Stateful DB, managed MQTT, production secrets/tenancy",
            ],
        ],
    )
    add_figure(
        doc,
        "fig-09-current-vs-target.png",
        "Figure 9. Solid boxes are implemented in these repos. Dashed boxes are the production target in docs/backend-architecture.md.",
    )

    add_heading(doc, "4. System context and repositories", 1)
    add_para(doc, "The product is two sibling repositories. The frontend never embeds backend source. The backend never re-creates the React app.")
    add_figure(
        doc,
        "fig-01-system-context.png",
        "Figure 1. Operators use the React console. Replay scripts and optional MQTT feed the FastAPI process. SQLite, model artifacts, and knowledge files stay on the backend.",
    )
    add_table(
        doc,
        ["Repository", "Role", "Primary stack"],
        [
            ["auravance-backend", "API, ingest, SQLite, models, RAG, Terraform, Docker", "Python 3.11+, FastAPI, scikit-learn, XGBoost, optional PyTorch"],
            ["auravance-frontend", "Operator console", "React 18, Vite 5, single src/App.jsx + styles.css"],
        ],
    )
    add_heading(doc, "4.1 Backend layout", 2)
    add_bullets(
        doc,
        [
            "app/main.py — FastAPI app, CORS, lifespan (SQLite init, MQTT start/stop)",
            "app/api/routes/ — REST surface under /api/v1",
            "app/models/schemas.py — Pydantic contracts (change these first)",
            "app/services/ — telemetry, scoring, anomaly, prognostics, sequence LSTM, RAG, alerts, master data, auth",
            "app/ingestion/ — replay helpers, MQTT protocol, embedded broker, subscriber",
            "app/core/ — settings, SQLite, JWT and ingest-key security",
            "scripts/ — ETL, training, replay, simulator, knowledge index",
            "artifacts/models/ — joblib artifacts + JSON sidecars",
            "knowledge/ — markdown guides and processed PDF chunks",
            "terraform/ — Cloud Run + Firebase Hosting site + secrets",
        ],
    )
    add_heading(doc, "4.2 Frontend layout", 2)
    add_bullets(
        doc,
        [
            "src/App.jsx — login, nav, Monitoring, Root Cause Analysis, master data, execution, admin",
            "src/styles.css — console styling",
            "vite.config.js — dev server on 5173, proxy /api → http://127.0.0.1:8000",
            "VITE_API_BASE_URL — optional; local default is /api/v1 through the proxy",
            "Session key localStorage: auravance.currentUser",
        ],
    )

    add_heading(doc, "5. Runtime architecture (what runs locally)", 1)
    add_para(
        doc,
        "Local topology: replay or simulator produces telemetry → REST ingest and/or MQTT → "
        "FastAPI telemetry_service persists to SQLite and scores the asset → alerts with cooldown → "
        "the React dashboard reads assets, telemetry history, predictions, alerts, and RAG over /api/v1.",
    )
    add_figure(
        doc,
        "fig-02-local-runtime.png",
        "Figure 2. What runs on a laptop: uvicorn :8000, Vite :5173, SQLite, optional embedded MQTT on 1883.",
    )
    add_para(
        doc,
        "simulate_telemetry.py / replay_telemetry.py\n"
        "        |  REST POST /api/v1/telemetry/ingest[/batch]\n"
        "        |  or MQTT plants/{plant}/lines/{line}/assets/{asset}/telemetry\n"
        "        v\n"
        "FastAPI (app/main.py)\n"
        "        |-- SQLite  data/auravance_demo.db\n"
        "        |-- score_asset()  Isolation Forest + XGBoost/HGB + LSTM (+ rules fallback)\n"
        "        |-- alerts (cooldown, default 300s)\n"
        "        +-- rag_service.root_cause()  knowledge chunks + OpenAI or playbook\n"
        "                ^\n"
        "auravance-frontend  src/App.jsx  (Vite proxy /api → :8000)",
        italic=True,
    )

    add_heading(doc, "6. Frontend application", 1)
    add_para(
        doc,
        "The dashboard is a React 18 single-page app. There is no client-side router library; "
        "navigation is module state in App.jsx. After JWT login, the console loads assets, plants, "
        "vendors, equipment, schedules, work orders, defects, users (admin), diagnostics status, "
        "then polls telemetry history and the latest prediction for the selected asset.",
    )
    add_figure(
        doc,
        "fig-03-frontend-modules.png",
        "Figure 3. Console modules after login. Operations read live telemetry and predictions; master data and execution are CRUD against SQLite.",
    )
    add_heading(doc, "6.1 Navigation groups", 2)
    add_table(
        doc,
        ["Group", "Modules", "Backend calls"],
        [
            ["Operations", "Monitoring; Root Cause Analysis", "telemetry history, predictions, alerts; POST /diagnostics/root-cause"],
            ["Master Data", "Plant Info; Equipment Master; Vendor Management", "CRUD /plants, /equipment, /vendors"],
            ["Execution", "Maintenance Schedules; Work Order Mgmt; Defects View", "CRUD /maintenance-schedules, /work-orders, /defects"],
            ["Admin", "User Mgmt (Admin role only)", "/users"],
        ],
    )
    add_heading(doc, "6.2 Monitoring and prediction UI", 2)
    add_para(
        doc,
        "Monitoring shows the four public sensor channels (temperature_c, vibration_rms, "
        "pressure_bar, current_a), trend charts from the last 24 history points, anomaly label, "
        "72-hour failure probability, estimated RUL hours, dominant signal, and alerts. Replay "
        "ground truth in source_context (NORMAL, BROKEN, RECOVERING, AIR_LEAK, FAILED) is shown "
        "next to model output when present.",
    )
    add_heading(doc, "6.3 How the frontend reaches the API", 2)
    add_bullets(
        doc,
        [
            "Local: fetch('/api/v1/...') → Vite proxy → http://127.0.0.1:8000/api/v1/...",
            "Deployed: VITE_API_BASE_URL=https://<cloud-run-host>/api/v1 baked in at npm run build",
            "Auth header: Authorization: Bearer <JWT> after POST /auth/login",
            "CORS on the API allows any origin (demo); cookies are not used",
        ],
    )
    add_figure(
        doc,
        "fig-10-request-path.png",
        "Figure 10. After login the SPA loads catalogs, then polls history and the latest prediction. RCA is an explicit POST, not part of ingest.",
    )

    add_heading(doc, "7. Backend application", 1)
    add_para(
        doc,
        "FastAPI exposes a versioned HTTP API at /api/v1. Lifespan initializes SQLite (DDL + demo seed) "
        "and starts the MQTT subscriber if AURAVANCE_MQTT_ENABLED is true. Model load failures must not "
        "crash the process: score_asset falls back to rules and status endpoints report fallback_active.",
    )
    add_heading(doc, "7.1 REST surface", 2)
    add_table(
        doc,
        ["Area", "Routes"],
        [
            ["Health", "GET /health (mounted at /api/v1/health)"],
            ["Auth", "POST /auth/login, GET /auth/me"],
            ["Assets", "GET/POST /assets"],
            ["Master data", "/plants, /vendors, /equipment, /sensors, /users, /maintenance-schedules, /work-orders, /defects"],
            ["Telemetry", "POST /telemetry/ingest, POST /telemetry/ingest/batch, GET /telemetry/ingest/targets, GET /telemetry/ingest/status, latest/history"],
            ["Predictions", "GET /predictions/assets/{asset_id}/latest, GET /predictions/models/status"],
            ["Diagnostics", "GET /diagnostics/status, POST /diagnostics/root-cause, GET /diagnostics/retrieve"],
            ["Knowledge", "/knowledge/documents, /knowledge/search"],
            ["Alerts", "GET /alerts"],
            ["Datasets", "GET /datasets"],
        ],
    )
    add_heading(doc, "7.2 Authentication split", 2)
    add_bullets(
        doc,
        [
            "Dashboard reads/writes: JWT Bearer from POST /auth/login. Demo admin is seeded in SQLite.",
            "Ingest: X-Ingest-Key when AURAVANCE_INGEST_API_KEY is set. Unset keeps ingest open for local development.",
            "These are different controls. A dashboard JWT cannot replace the ingest key, and vice versa.",
        ],
    )

    add_heading(doc, "8. Ingestion pipelines", 1)
    add_heading(doc, "8.1 Live ingest contract", 2)
    add_para(
        doc,
        "Every producer — REST simulator, REST replay, MQTT subscriber — must end in telemetry_service. "
        "Scoring and alerting are not bypassed. The public payload is:",
    )
    add_para(
        doc,
        '{ "timestamp": "2026-05-15T14:22:31Z", "asset_id": "compressor-17", '
        '"metrics": { "temperature_c": 84.2, "vibration_rms": 0.63, "pressure_bar": 5.2, "current_a": 11.7 }, '
        '"source": "api", "source_context": {} }',
        italic=True,
    )
    add_bullets(
        doc,
        [
            "Four channels only on the wire. Expanding them requires the frontend and all three model profiles together.",
            "Batch: {\"readings\":[...]} min 1, max 500. Unknown asset IDs fail the whole batch (404).",
            "Replay fills source_context with original timestamp, raw values, ground-truth status, units, and dataset id.",
            "On ingest, the last 8 points are scored for alerts. Cooldown AURAVANCE_ALERT_COOLDOWN_SECONDS (default 300) unless severity changes.",
        ],
    )
    add_figure(
        doc,
        "fig-04-ingestion.png",
        "Figure 4. REST and MQTT are producers. Scoring, SQLite writes, and alert cooldown all happen inside telemetry_service.",
    )
    add_heading(doc, "8.2 REST ingest", 2)
    add_para(
        doc,
        "POST /api/v1/telemetry/ingest and /ingest/batch are the original and still-supported path. "
        "scripts/simulate_telemetry.py posts synthetic readings. scripts/replay_telemetry.py "
        "(--transport rest, default) posts one batch per tick for every registered asset.",
    )
    add_heading(doc, "8.3 MQTT ingest (local)", 2)
    add_para(
        doc,
        "MQTT is an additional producer into the same service. Default local settings start an embedded "
        "MQTT 3.1.1 broker (QoS 0) and a paho-mqtt subscriber. Topic:",
    )
    add_para(doc, "plants/{plant_id}/lines/{line_id}/assets/{asset_id}/telemetry", italic=True)
    add_bullets(
        doc,
        [
            "Subscriber: app/ingestion/mqtt.py — decode JSON, map topic asset_id, call telemetry_service.ingest",
            "Status: GET /telemetry/ingest/status (ingest-key gated)",
            "Replay: python scripts/replay_telemetry.py --transport mqtt",
            "uvicorn --reload restarts the process and drops the embedded broker; REST replay is more reliable while developing",
            "Production still needs an external broker. The in-process broker is a local stand-in, not a plant MQTT cluster",
        ],
    )
    add_heading(doc, "8.4 Offline ETL and replay datasets", 2)
    add_para(
        doc,
        "High-volume plant ETL is not a streaming job in this repo. Offline scripts prepare public datasets "
        "used for training and for honest live replay.",
    )
    add_table(
        doc,
        ["Profile", "Asset types", "Source dataset", "Sampling", "Failure character"],
        [
            ["pump", "pump, cooling-system", "Kaggle pump-sensor (replay file committed; full CSV gitignored)", "1 min", "Mixed; 72h forecast weak"],
            ["compressor", "compressor", "MetroPT-3 air compressor, 4 air leaks", "5 min", "Sudden; leak-in-progress is learnable, 72h is not"],
            ["motor", "motor, drive, conveyor", "NASA/IMS bearing run-to-failure", "10 min vibration RMS", "Gradual degradation"],
        ],
    )
    add_bullets(
        doc,
        [
            "scripts/prepare_datasets.py, prepare_metropt.py, prepare_ims.py, build_replay_dataset.py",
            "Replay files in data/replay/ are what local demo streaming uses",
            "GET /datasets lists the catalog stored in SQLite",
            "This is not a 1M+ records/day pipeline. That remains a TimescaleDB-era target.",
        ],
    )

    add_heading(doc, "9. Databases and storage", 1)
    add_heading(doc, "9.1 SQLite demo store (implemented)", 2)
    add_para(
        doc,
        "AURAVANCE_DEMO_DB_PATH defaults to data/auravance_demo.db. init_db() creates tables and seeds "
        "demo plants, assets, vendors, equipment, users, schedules, work orders, and defects. Schema "
        "changes belong in app/core/db.py together with seed data. WAL mode is enabled for concurrent "
        "API + MQTT threads.",
    )
    add_table(
        doc,
        ["Table group", "Tables", "Purpose"],
        [
            ["Telemetry path", "assets, telemetry, alerts", "Live readings and scored warnings"],
            ["Master data", "plants, vendors, equipment, sensors, users", "Dashboard CRUD"],
            ["Execution", "maintenance_schedules, work_orders, defects", "Seeded workflows; not auto-opened from alerts"],
            ["Catalog / RAG", "dataset_catalog, knowledge documents/chunks/vectors", "Dataset registry and retrieval index"],
        ],
    )
    add_para(
        doc,
        "Telemetry columns: asset_id, timestamp, temperature_c, vibration_rms, pressure_bar, current_a, "
        "source, source_context_json. Index: (asset_id, timestamp).",
    )
    add_figure(
        doc,
        "fig-07-sqlite.png",
        "Figure 7. Implemented tables in SQLite. Dashed boxes are the TimescaleDB / pgvector cutover, which is not in this repo.",
    )
    add_heading(doc, "9.2 What SQLite is not", 2)
    add_bullets(
        doc,
        [
            "Not TimescaleDB. No hypertables, no continuous aggregates, no retention policies.",
            "Not pgvector. Embeddings live as JSON/blobs in SQLite when OPENAI_API_KEY is set during indexing.",
            "Not durable on Cloud Run. The container filesystem is ephemeral; a new revision starts from an empty seeded DB unless a volume is added (not in Terraform today).",
        ],
    )
    add_heading(doc, "9.3 Production storage target", 2)
    add_para(
        doc,
        "docs/backend-architecture.md specifies PostgreSQL + TimescaleDB for telemetry and metadata, "
        "and pgvector for RAG. That cutover is not implemented. Object storage for manuals and model "
        "artifacts is also a target (S3-compatible), while today artifacts and knowledge are copied into the image.",
    )

    add_heading(doc, "10. Machine learning and scoring", 1)
    add_para(
        doc,
        "Three tasks stay separate: anomaly (unsupervised Isolation Forest on healthy operation), "
        "failure (P(failing now or within 72h)), and RUL (hours until next failure, capped). "
        "Do not reuse a classification score as RUL except the documented compressor sudden-failure wrappers.",
    )
    add_heading(doc, "10.1 Serving path", 2)
    add_bullets(
        doc,
        [
            "score_asset in app/services/scoring.py",
            "Profile from asset type: pump / compressor / motor (cooling-system → pump; drive/conveyor → motor)",
            "Anomaly: Isolation Forest if loaded, else rules mix of normalized channels + trend",
            "Failure: LSTM if history length ≥ window (12) and artifact loaded; else XGBoost if present; else HGB; else rules",
            "RUL: XGBoost if present, else HGB, else derived from failure probability",
            "Predictions use last 12 points; ingest alerts use last 8",
            "Honest replay: if source_context.dataset matches fold_dataset and the timestamp is in a held-out window, serve +heldout-N",
        ],
    )
    add_figure(
        doc,
        "fig-05-scoring.png",
        "Figure 5. Isolation Forest for anomaly, LSTM then XGBoost/HGB for failure, XGBoost/HGB for RUL, rules if artifacts fail to load.",
    )
    add_heading(doc, "10.2 Artifacts", 2)
    add_para(
        doc,
        "joblib dict, artifact_schema_version: 1, plus a sibling .json with the estimator stripped. "
        "GET /predictions/models/status reads metadata (algorithm, metrics, fallback_active) per profile, "
        "including sequence/LSTM status.",
    )
    add_heading(doc, "10.3 Accuracy honesty", 2)
    add_para(
        doc,
        "Published leave-one-failure-out numbers live in artifact JSON (pump failure AUC ~0.66 on the "
        "full HGB training set; compressor leak-now is strong and 72h forecast is not; motor bearings degrade). "
        "scripts/evaluate_detection_delta.py compares models to the rules scorer on replay. Pump beats a weak "
        "rules baseline by a large relative margin on that clip-based detection task; compressor is about +20%; "
        "motor snapshots are too sparse for a 30% claim. Do not state a product-wide 30% detection-accuracy improvement.",
    )

    add_heading(doc, "11. RAG diagnostics", 1)
    add_para(
        doc,
        "POST /diagnostics/root-cause takes asset_id, symptom_summary, and optional alert ids. The service "
        "loads the latest prediction, maps dominant_signal to playbook terms, retrieves knowledge chunks "
        "(OpenAI embeddings when the key is present, else TF-IDF), then calls the OpenAI Responses API with "
        "a JSON schema. If the key or API is down, it extracts cause/action sentences from chunks plus the playbook.",
    )
    add_figure(
        doc,
        "fig-06-rag.png",
        "Figure 6. RCA is request-driven. Retrieval uses OpenAI embeddings or TF-IDF; generation uses the Responses API or the local playbook.",
    )
    add_bullets(
        doc,
        [
            "Index: python scripts/prepare_knowledge.py over knowledge/raw PDFs and knowledge/*.md",
            "Return fields: grounded_summary, likely_causes, recommended_actions, confidence, generation_mode, generation_model, evidence_snippets",
            "Status reports openai_api_key_loaded as a boolean only — the key is never echoed",
            "Diagnostics are request-driven, not on the ingest hot path",
            "pgvector is the documented next index, with the same retrieve() signature",
        ],
    )

    add_heading(doc, "12. Local development process", 1)
    add_para(doc, "Python 3.11 or 3.12 is required. Python 3.14 may not have XGBoost/PyTorch wheels.")
    add_para(
        doc,
        "Backend:\n"
        "  py -3.12 -m venv .venv\n"
        "  .venv/Scripts/python.exe -m pip install -e \".[sequence]\"\n"
        "  .venv/Scripts/python.exe -m uvicorn app.main:app --reload --host 127.0.0.1 --port 8000\n\n"
        "Frontend (sibling repo):\n"
        "  npm install && npm run dev     # http://localhost:5173\n\n"
        "Optional live replay:\n"
        "  .venv/Scripts/python.exe scripts/replay_telemetry.py --transport rest",
        italic=True,
    )
    add_para(doc, "Copy .env.example to .env for OPENAI_API_KEY, ingest key, MQTT flags, and auth secret. Never commit .env.")

    add_heading(doc, "13. Docker", 1)
    add_para(
        doc,
        "auravance-backend/Dockerfile builds a python:3.11-slim image. It installs build-essential, copies "
        "pyproject.toml, app/, artifacts/, and knowledge/, then pip install --no-cache-dir -e . "
        "The default extra does not include PyTorch, so Cloud Run serves Isolation Forest + XGBoost/HGB "
        "and LSTM falls back unless the image is built with the sequence extra.",
    )
    add_bullets(
        doc,
        [
            "CMD: uvicorn app.main:app --host 0.0.0.0 --port ${PORT}  (Cloud Run injects PORT; Dockerfile default 8080)",
            "data/, scripts/, and the SQLite file are not copied — replay ETL is a local/dev concern",
            "Model artifacts in artifacts/models are baked into the image",
            "There is no docker-compose for Postgres or Mosquitto in this tree",
        ],
    )
    add_para(
        doc,
        "docker build -t us-central1-docker.pkg.dev/<project>/auravance/backend:latest .\n"
        "docker push us-central1-docker.pkg.dev/<project>/auravance/backend:latest",
        italic=True,
    )

    add_heading(doc, "14. Terraform and GCP demo deploy", 1)
    add_para(
        doc,
        "terraform/ is a free-tier portfolio deploy, not the MQTT/Timescale production design. "
        "Estimated cost at demo traffic is intended to stay inside GCP Always Free limits "
        "(Cloud Run scale-to-zero, 2M requests/month class).",
    )
    add_figure(
        doc,
        "fig-08-deploy.png",
        "Figure 8. Docker image → Artifact Registry → Cloud Run. Terraform creates the site resource; firebase deploy uploads the SPA.",
    )
    add_heading(doc, "14.1 Resources", 2)
    add_table(
        doc,
        ["File", "What it creates"],
        [
            ["apis.tf", "Enables run, artifactregistry, secretmanager, firebase, firebasehosting APIs"],
            ["artifact_registry.tf", "Docker repo auravance in the chosen region"],
            ["secrets.tf", "Secret Manager openai-api-key; dedicated Cloud Run service account with secretAccessor"],
            ["cloud_run.tf", "Cloud Run v2 service, 1 vCPU / 512Mi, min instances 0, max 2, port 8080, public invoker optional"],
            ["firebase_hosting.tf", "Firebase project + Hosting site resource; it does not upload the SPA files"],
            ["variables.tf", "project_id, region (us-central1), backend_image, openai_api_key (sensitive), allow_unauthenticated"],
        ],
    )
    add_heading(doc, "14.2 Deploy sequence", 2)
    add_bullets(
        doc,
        [
            "1. Create/select a GCP project and enable billing (required even for free tier).",
            "2. Build and push the backend image to Artifact Registry (or apply the registry resource first).",
            "3. terraform init && terraform apply with terraform.tfvars (gitignored). Never commit terraform.tfvars.",
            "4. terraform output backend_url",
            "5. In auravance-frontend: set VITE_API_BASE_URL to that URL plus /api/v1, npm run build.",
            "6. firebase init hosting (public dir dist) and firebase deploy --only hosting.",
            "7. Health check GET {backend_url}/api/v1/health — not /health at the container root.",
        ],
    )
    add_heading(doc, "14.3 What Terraform does not provision", 2)
    add_bullets(
        doc,
        [
            "No Cloud SQL / TimescaleDB",
            "No Memorystore, no MQTT broker (HiveMQ, Mosquitto, Cloud IoT)",
            "No load balancer beyond Cloud Run URL",
            "No CI pipeline; image build and firebase deploy are manual CLI steps",
            "SQLite inside Cloud Run is ephemeral — demo data resets when the instance is replaced",
        ],
    )

    add_heading(doc, "15. Security notes", 1)
    add_bullets(
        doc,
        [
            "Secrets from environment / .env / Secret Manager only. Do not commit .env or terraform.tfvars.",
            "JWT secret AURAVANCE_AUTH_SECRET_KEY must be changed off the demo default before any shared deploy.",
            "Ingest key is optional locally and should be set in any network-exposed API.",
            "Cloud Run allow_unauthenticated defaults true for a portfolio demo; that is not a plant-network posture.",
            "RAG status exposes openai_api_key_loaded as boolean only.",
            "Password hashes are stored in SQLite users.password_hash.",
        ],
    )

    add_heading(doc, "16. Production target architecture (not built)", 1)
    add_para(
        doc,
        "The longer-term backend design in docs/backend-architecture.md is: industrial sensors/PLCs → "
        "MQTT broker → ingestion service → stream processing → TimescaleDB and a feature store → "
        "XGBoost/LSTM/Isolation Forest inference → alerts and maintenance workflows → RAG over pgvector. "
        "The FastAPI gateway remains the dashboard API. Extra production pieces named there and not present "
        "in code: Celery/Dramatiq, Redis/RabbitMQ, WebSocket or SSE push, S3-compatible object storage, "
        "Kubernetes, and auto work-order generation from alerts.",
    )
    add_para(
        doc,
        "Documented upgrade order: keep REST ingest; add external MQTT into telemetry_service; serve XGBoost "
        "behind artifact schema v1 (done locally); add LSTM beside tabular (done locally); move telemetry to "
        "TimescaleDB when volume needs it; measure detection accuracy against the rules baseline before claiming 30%.",
    )

    add_heading(doc, "17. HTTP contracts to keep stable", 1)
    add_para(
        doc,
        "The frontend depends on these staying compatible. Additive fields are acceptable; renaming or removing "
        "prediction/telemetry/RAG fields requires a lockstep change in auravance-frontend.",
    )
    add_table(
        doc,
        ["Contract", "Stability rule"],
        [
            ["Telemetry ingest", "Four metrics keys; source_context preserved"],
            ["Prediction latest", "anomaly_score/label, failure_probability_72h, estimated_rul_hours, dominant_signal, model versions"],
            ["RAG root-cause", "likely_causes, recommended_actions, grounded_summary, evidence_snippets, generation_mode"],
            ["Auth", "POST /auth/login returns access_token; subsequent calls use Bearer"],
        ],
    )

    add_heading(doc, "18. Related files", 1)
    add_bullets(
        doc,
        [
            "auravance-backend/docs/Auravance-Architecture.docx — this illustrated write-up",
            "auravance-backend/docs/architecture-diagrams/ — PNG sources for the figures",
            "auravance-backend/docs/demo-backend-architecture.md — local stack",
            "auravance-backend/docs/backend-architecture.md — production target",
            "auravance-backend/terraform/README-DEPLOY.md — GCP steps",
            "auravance-backend/README.md — run, train, replay",
            "auravance-frontend/README.md — Vite proxy and production API URL",
        ],
    )

    footer = doc.add_paragraph()
    footer.paragraph_format.space_before = Pt(18)
    run = footer.add_run(
        "End of document. This architecture write-up is generated from the repository layout as of 21 October 2026 "
        "and is intended for engineering onboarding, not as a resume substitute."
    )
    set_run(run, italic=True, size=10, color=MUTED)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    path = build()
    print(path)
