"""Generate docs/Auravance-Data-Workflow.docx from the live backend, not the resume slide."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BACKEND_DIR = Path(__file__).resolve().parents[1]
OUTPUT = BACKEND_DIR / "docs" / "Auravance-Data-Workflow.docx"
FIGURES = BACKEND_DIR / "docs" / "architecture-diagrams"

NAVY = RGBColor(0x0B, 0x2E, 0x4A)
TEAL = RGBColor(0x0D, 0x6E, 0x7A)
MUTED = RGBColor(0x44, 0x55, 0x66)


def set_run(run, *, bold=False, size=11, color=None, italic=False):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    if color is not None:
        run.font.color.rgb = color


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = NAVY if level == 1 else TEAL
        run.font.name = "Calibri"
    return p


def add_para(doc, text, *, bold=False, italic=False, size=11, space_after=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run(text)
    set_run(run, bold=bold, italic=italic, size=size, color=MUTED if italic and not bold else None)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(item, style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        for run in p.runs:
            run.font.name = "Calibri"
            run.font.size = Pt(11)


def add_figure(doc, filename: str, caption: str) -> None:
    path = FIGURES / filename
    if not path.exists():
        raise FileNotFoundError(f"Missing diagram {path}. Run scripts/generate_architecture_diagrams.py first.")
    picture = doc.add_paragraph()
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.paragraph_format.space_before = Pt(8)
    picture.paragraph_format.space_after = Pt(2)
    run = picture.add_run()
    run.add_picture(str(path), width=Inches(6.5))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(12)
    r = cap.add_run(caption)
    set_run(r, italic=True, size=10, color=MUTED)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                set_run(run, bold=True, size=10, color=NAVY)
    for r_index, row in enumerate(rows, start=1):
        for c_index, value in enumerate(row):
            cell = table.rows[r_index].cells[c_index]
            cell.text = str(value)
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    set_run(run, size=10)
    doc.add_paragraph()
    return table


def build() -> Path:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.9)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    title = doc.add_paragraph()
    run = title.add_run("Auravance")
    set_run(run, bold=True, size=28, color=NAVY)

    subtitle = doc.add_paragraph()
    run = subtitle.add_run("Data workflow: extract, load, store, and stream")
    set_run(run, size=16, color=TEAL)

    meta = doc.add_paragraph()
    run = meta.add_run(
        "This document describes how sensor data moves through the repositories "
        "as they exist today. Production targets such as 1M+ records/day, an HA MQTT "
        "cluster, pgvector, and a measured 30% detection improvement are called out "
        "separately and are not treated as live behavior."
    )
    set_run(run, size=11, italic=True, color=MUTED)

    add_para(
        doc,
        "Repositories: auravance-backend (FastAPI) and auravance-frontend (React 18 + Vite). "
        "Date: 26 October 2026.",
    )

    add_heading(doc, "1. End-to-end workflow", 1)
    add_para(
        doc,
        "Auravance is a predictive-maintenance console. Recorded industrial datasets "
        "are prepared offline, then replayed as if they were live plant telemetry. "
        "The FastAPI service accepts those readings, writes them to a database, "
        "scores failure risk, and serves history, predictions, alerts, and diagnostics "
        "to the React dashboard over HTTP /api/v1.",
    )
    add_para(doc, "The live path is:")
    add_bullets(
        doc,
        [
            "Extract: offline scripts turn public research datasets into replay CSVs and training tables.",
            "Load: init_db() seeds plants, assets, and master data; ingest writes sensor rows into telemetry.",
            "Stream: replay_telemetry.py or simulate_telemetry.py push readings over REST or MQTT.",
            "Store: PostgreSQL + Timescale when AURAVANCE_DATABASE_URL is set; otherwise SQLite WAL.",
            "Score: Isolation Forest plus XGBoost (or HistGradientBoosting) and optional LSTM, with a rules fallback.",
            "Read: the dashboard polls history, latest prediction, and alerts for the selected asset.",
        ],
    )
    add_figure(doc, "fig-01-system-context.png", "Figure 1. Two repositories, one HTTP contract.")
    add_figure(doc, "fig-02-local-runtime.png", "Figure 2. Local processes: API, replay, optional Docker stores.")

    add_heading(doc, "2. Where the data comes from", 1)
    add_para(
        doc,
        "This demo does not attach to a live plant historian. Streaming comes from "
        "three published failure recordings that already sit under data/replay/. "
        "Each registered asset is mapped to the dataset that matches its equipment profile.",
    )
    add_table(
        doc,
        ["Asset types", "Profile", "Source dataset", "Native sampling", "What failed in the original study"],
        [
            [
                "pump, cooling-system",
                "pump",
                "Kaggle pump-sensor recording",
                "1 minute",
                "7 breakdowns (anonymized channels)",
            ],
            [
                "compressor",
                "compressor",
                "MetroPT-3 metro air compressor (UCI, CC BY 4.0)",
                "5-minute averages",
                "4 air leaks from the operator report",
            ],
            [
                "motor, drive, conveyor",
                "motor",
                "NASA/IMS bearing run-to-failure (AC motor at 2000 rpm)",
                "1 s vibration snapshot every 10 minutes",
                "3 bearing failures",
            ],
        ],
    )
    add_para(
        doc,
        "Default demo assets are seeded at API startup (pump-04, compressor-17, motor-08, "
        "and the rest of the ten-asset plant). Replay discovers them through "
        "GET /api/v1/telemetry/ingest/targets and re-checks that list every 60 seconds, "
        "so new assets join the stream without a restart.",
    )
    add_para(
        doc,
        "Enterprise plants would publish the same JSON onto MQTT from gateways or "
        "historians. That producer is not in this repo; compose Mosquitto is the local "
        "broker stand-in. A multi-node EMQX/HiveMQ cluster remains an operator-owned target.",
        italic=True,
    )
    add_figure(
        doc,
        "fig-12-demo-simulation.png",
        "Figure 3. Demo path: recorded episodes replayed as live telemetry.",
    )
    add_figure(
        doc,
        "fig-11-enterprise-sensors.png",
        "Figure 4. Target plant path: sensors to MQTT to the same ingest contract.",
    )

    add_heading(doc, "3. Extract: how raw files become contract metrics", 1)
    add_para(
        doc,
        "Extraction is offline. Operators place original archives under data/raw/ "
        "(and related interim folders). Dedicated scripts validate, unpack, resample, "
        "and write normalized CSVs. Those files are the training and replay inputs; "
        "the API never reads the original Kaggle, UCI, or IMS archives at request time.",
    )
    add_table(
        doc,
        ["Script", "Input", "Output"],
        [
            [
                "scripts/prepare_datasets.py",
                "AI4I, Kaggle pump, C-MAPSS zip, IMS 7z under data/raw/",
                "Normalized CSVs in data/processed plus dataset registry rows",
            ],
            [
                "scripts/build_replay_dataset.py",
                "Processed Kaggle pump CSV",
                "data/replay/pump_failure_episodes.csv",
            ],
            [
                "scripts/prepare_metropt.py",
                "MetroPT-3 zip in data/raw/MetroPT3/",
                "Compressor replay CSV (oil temperature, panel pressure, motor current)",
            ],
            [
                "scripts/prepare_ims.py",
                "IMS .rar extracts (7-Zip); caches features per test",
                "Motor replay CSV (worst-bearing vibration RMS in g)",
            ],
        ],
    )
    add_para(
        doc,
        "On the wire every reading uses four named channels: temperature_c, vibration_rms, "
        "pressure_bar, and current_a. Replay maps each dataset onto that contract and keeps "
        "the original timestamp, raw values, units, dataset id, and ground-truth machine "
        "status in source_context. The dashboard shows that ground truth next to the model.",
    )
    add_bullets(
        doc,
        [
            "Pump: anonymized Kaggle channels have unknown units. Replay rescales them with "
            "the Isolation Forest sidecar (median / IQR into the model's feature space) so "
            "serving sees the same four features used in training.",
            "Compressor: oil_temperature_c → temperature_c, tp3_bar → pressure_bar, "
            "motor_current_a → current_a. Vibration is not present in the MetroPT-3 mapping.",
            "Motor: max_bearing_rms → vibration_rms (g). Temperature, pressure, and current "
            "are not in the IMS mapping.",
            "Ground-truth labels travel as NORMAL, BROKEN, RECOVERING, AIR_LEAK, or FAILED.",
        ],
    )
    add_para(
        doc,
        "Training is a separate extract-and-fit path. scripts/train_xgboost_models.py, "
        "scripts/train_sequence_models.py, and related trainers write joblib artifacts "
        "plus sidecar JSON under artifacts/models/. Serving loads those files; it does "
        "not retrain on ingest.",
    )

    add_heading(doc, "4. Load: how readings enter the platform", 1)
    add_para(
        doc,
        "There are two load stages. First, init_db() creates tables and seeds demo "
        "assets, plants, vendors, equipment, users, and related master data. Second, "
        "every live reading is loaded by TelemetryService.ingest or ingest_batch.",
    )
    add_para(doc, "Ingest payload (single reading):")
    add_para(
        doc,
        '{ "timestamp": "2026-05-15T14:22:31Z", "asset_id": "compressor-17", '
        '"metrics": { "temperature_c": 84.2, "vibration_rms": 0.63, '
        '"pressure_bar": 5.2, "current_a": 11.7 }, "source": "api", "source_context": {} }',
        italic=True,
    )
    add_bullets(
        doc,
        [
            "POST /api/v1/telemetry/ingest accepts one reading.",
            "POST /api/v1/telemetry/ingest/batch accepts {\"readings\": [...]} with 1 to 500 rows. "
            "Unknown asset ids fail the whole batch (404).",
            "MQTT messages on plants/{plant_id}/lines/{line_id}/assets/{asset_id}/telemetry "
            "are decoded and forwarded into the same telemetry_service.ingest path. "
            "Scoring is not skipped.",
            "When AURAVANCE_INGEST_API_KEY is set, ingest requires X-Ingest-Key. "
            "Dashboard reads use JWT Bearer instead. Local demo often leaves the ingest key empty.",
        ],
    )
    add_para(
        doc,
        "Each accepted reading is inserted in one transaction into telemetry "
        "(asset_id, timestamp, the four channels, source, source_context_json). "
        "The API then schedules scoring for that asset: a Redis list push when "
        "AURAVANCE_REDIS_URL is set, otherwise score_ingest_alerts() runs inline "
        "in the API process from the last eight points.",
    )
    add_figure(doc, "fig-04-ingestion.png", "Figure 5. REST and MQTT converge on telemetry_service.")

    add_heading(doc, "5. How the data is streamed, and from where", 1)
    add_para(
        doc,
        "Streaming is a producer process outside the API. The default producer is "
        "scripts/replay_telemetry.py. It loads the three replay CSVs, asks the API "
        "which assets exist, and on each tick sends one batch with the next row for "
        "every asset. Default pace is five recorded minutes every two seconds. "
        "Timestamps on the wire are wall-clock now; the original recorded time stays "
        "in source_context. Assets that share a dataset start at staggered offsets "
        "so the plant is not all in the same failure state at once.",
    )
    add_para(
        doc,
        "Transport is REST by default (POST /telemetry/ingest/batch). Pass --transport mqtt "
        "to publish the same JSON onto Mosquitto. MQTT still uses REST only to discover "
        "assets. scripts/simulate_telemetry.py is the synthetic alternative: it fabricates "
        "the four channels instead of replaying a public dataset.",
    )
    add_para(
        doc,
        "If replay is not running, the dashboard still lists assets (seeded master data) "
        "but history, charts, and predictions stay empty. That is expected: the API does "
        "not generate telemetry by itself.",
    )
    add_para(doc, "MQTT broker choice:", bold=True)
    add_bullets(
        doc,
        [
            "AURAVANCE_MQTT_EMBEDDED=true (default): an in-process broker in the API, used for laptop and pytest.",
            "AURAVANCE_MQTT_EMBEDDED=false: the API subscribes to compose Mosquitto on 127.0.0.1:1883. "
            "If the broker is down, REST ingest still starts; MQTT is best-effort.",
            "Topic contract: plants/{plant_id}/lines/{line_id}/assets/{asset_id}/telemetry.",
        ],
    )

    add_heading(doc, "6. Where the data is stored", 1)
    add_para(
        doc,
        "Storage is selected at process start. app/core/database.py uses PostgreSQL "
        "when AURAVANCE_DATABASE_URL is a non-empty string; otherwise it opens SQLite. "
        "Callers keep ? placeholders and dict-like rows either way.",
    )
    add_table(
        doc,
        ["Store", "When it is used", "What lives there"],
        [
            [
                "PostgreSQL 16 + TimescaleDB",
                "AURAVANCE_DATABASE_URL set (docker compose publishes host port 5433)",
                "assets, telemetry (hypertable on timestamp), alerts, plants, vendors, "
                "equipment, sensors, users, work orders, defects, knowledge chunks, dataset registry",
            ],
            [
                "SQLite WAL",
                "URL unset: pytest, no-Docker laptop, AURAVANCE_DEMO_DB_PATH "
                "(default data/auravance_demo.db)",
                "The same relational schema via app/core/db.py",
            ],
            [
                "Redis 7",
                "AURAVANCE_REDIS_URL set (compose port 6379)",
                "List auravance:score of asset ids waiting for ingest scoring. Not a sensor store.",
            ],
            [
                "Local files",
                "Always",
                "data/replay/*.csv, data/processed, artifacts/models/*.joblib + JSON, knowledge/",
            ],
        ],
    )
    add_figure(
        doc,
        "fig-07-sqlite.png",
        "Figure 6. Relational demo schema. Timescale only changes how telemetry is physically stored.",
    )

    add_heading(doc, "7. Which database is used, and why", 1)
    add_heading(doc, "7.1 PostgreSQL + Timescale (platform path)", 2)
    add_para(
        doc,
        "When Docker Desktop is running and compose is up, the API talks to "
        "postgresql://auravance:auravance@127.0.0.1:5433/auravance. The container image "
        "is timescale/timescaledb:2.17.2-pg16. After ordinary CREATE TABLE statements, "
        "postgres_schema.py enables the timescaledb extension and calls create_hypertable "
        "on telemetry, partitioned on timestamp. If that call fails, the table stays a "
        "regular PostgreSQL table and the API still runs.",
    )
    add_para(doc, "Why PostgreSQL for this product:")
    add_bullets(
        doc,
        [
            "Relational integrity: telemetry and alerts reference assets; master data "
            "(plants, vendors, work orders) needs foreign keys and SQL that the dashboard already assumes.",
            "One process can serve both high-frequency sensor rows and slower CMMS-style records "
            "without a second application database.",
            "psycopg is already on the dependency list; the adapter keeps the service code dialect-light.",
            "The production target is still SQL (retention policies, more writers). Starting on Postgres "
            "avoids a later SQLite-to-Postgres rewrite of every query.",
        ],
    )
    add_para(doc, "Why Timescale on the telemetry table specifically:")
    add_bullets(
        doc,
        [
            "Telemetry is time-series: many rows per asset_id, always filtered by timestamp. "
            "A hypertable chunks by time so recent history queries stay cheap as volume grows.",
            "The north-star load (1M+ sensor records/day, sub-minute ingest-to-score latency) "
            "is the reason Timescale is in compose. The local demo does not yet enforce that volume "
            "or retention policies; those remain targets.",
            "Chunking by time is a better fit than a generic B-tree-only table once replay or a "
            "plant gateway writes continuously.",
        ],
    )

    add_heading(doc, "7.2 SQLite WAL (laptop and tests)", 2)
    add_para(
        doc,
        "SQLite is the default when AURAVANCE_DATABASE_URL is blank. It is a single file "
        "with WAL journaling and foreign keys on. Pytest and a no-Docker laptop use it so "
        "the API starts without Docker Desktop, without a listening port 5433, and without "
        "a Postgres password.",
    )
    add_para(doc, "Why SQLite is kept:")
    add_bullets(
        doc,
        [
            "Zero operational dependencies for unit tests and first-run demos.",
            "Same logical schema as Postgres, so services are not forked into two code paths.",
            "WAL allows the API and a reader (or replay, via HTTP) to coexist on one file.",
        ],
    )
    add_para(
        doc,
        "SQLite is not the plant-scale store. Concurrent writers, retention, and "
        "time-based partitioning are why the Docker path exists.",
        italic=True,
    )

    add_heading(doc, "7.3 Redis (scoring queue, not telemetry)", 2)
    add_para(
        doc,
        "Redis is optional. With AURAVANCE_REDIS_URL set, ingest only LPUSHes the asset_id "
        "onto auravance:score. app.workers.scoring_worker pops that list and runs "
        "score_ingest_alerts. If Redis is down, ingest falls back to inline scoring so "
        "alerts are not silently dropped. Redis is used because ingest should not block "
        "the HTTP write path on model inference; it is not used as the system of record.",
    )

    add_heading(doc, "8. After storage: scoring, alerts, and the dashboard", 1)
    add_para(
        doc,
        "GET /api/v1/predictions/assets/{asset_id}/latest reads the last 12 telemetry "
        "points and calls score_asset. Ingest alerts use the last 8 points, with a cooldown "
        "(default 300 seconds) unless severity changes. Models are Isolation Forest for "
        "anomaly, XGBoost tabular failure and RUL when artifacts exist (otherwise "
        "HistGradientBoosting), and LSTM sequence scoring beside tabular when the optional "
        "sequence extra and artifacts are present. If models fail to load, the rules scorer "
        "runs and the response reports fallback_active.",
    )
    add_para(
        doc,
        "Replay rows whose source_context.dataset matches the training set use the stored "
        "leave-one-failure-out fold estimator (+heldout-N) so the demo does not score a "
        "failure with an estimator that already trained on that same episode.",
    )
    add_figure(doc, "fig-05-scoring.png", "Figure 7. Scoring path after telemetry is stored.")
    add_para(
        doc,
        "The React console (localhost:5173) proxies /api to the API on port 8000. After login "
        "it lists assets, then polls telemetry history and the latest prediction for the "
        "selected machine. Charts stay empty until ingest has written rows. RAG diagnostics "
        "read knowledge chunks (OpenAI embeddings or TF-IDF, playbook fallback). pgvector "
        "is the intended RAG index and is not the live store.",
    )

    add_heading(doc, "9. Local run order that actually fills the console", 1)
    add_bullets(
        doc,
        [
            "Start Docker Desktop, then docker compose up -d (Timescale, Mosquitto, Redis) if using the platform .env.",
            "Start the API: python -m uvicorn app.main:app --reload",
            "Start the scoring worker if Redis is set: python -m app.workers.scoring_worker",
            "Start the dashboard from auravance-frontend: npm run dev",
            "Start the producer: python scripts/replay_telemetry.py",
        ],
    )
    add_para(
        doc,
        "Without the last step, Postgres can be healthy and the dashboard can still show "
        "\"Waiting for more telemetry\". Seeded assets are not the same as sensor history.",
    )

    add_heading(doc, "10. What this workflow does not do yet", 1)
    add_bullets(
        doc,
        [
            "No live plant OPC-UA / historian connector in this repo.",
            "Compose Mosquitto is one broker, not an HA MQTT cluster.",
            "No automated retention or 1M+ records/day SLA in the local tree.",
            "No database backups in this stack.",
            "pgvector is not the live RAG index.",
            "Do not claim a 30% detection improvement until scripts/evaluate_detection_delta.py "
            "has published a measured baseline-versus-model delta.",
        ],
    )

    add_heading(doc, "Appendix. Environment switches that change this workflow", 1)
    add_table(
        doc,
        ["Variable", "Effect on data flow"],
        [
            ["AURAVANCE_DATABASE_URL", "Non-empty: PostgreSQL/Timescale. Empty: SQLite WAL."],
            ["AURAVANCE_REDIS_URL", "Non-empty: ingest scoring queued. Empty: inline scoring."],
            ["AURAVANCE_MQTT_EMBEDDED", "true: in-process broker. false: subscribe to compose Mosquitto."],
            ["AURAVANCE_MQTT_ENABLED", "false disables the MQTT subscriber; REST ingest still works."],
            ["AURAVANCE_INGEST_API_KEY", "When set, replay and MQTT must present X-Ingest-Key."],
            ["AURAVANCE_DEMO_DB_PATH", "SQLite file location when Postgres is not used."],
        ],
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    path = build()
    print(path)
